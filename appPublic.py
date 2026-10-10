from flask import Flask, request
import requests
from requests.auth import HTTPBasicAuth
import pytesseract
from PIL import Image
import cv2
import time
from datetime import datetime
import re 
from openpyxl import Workbook, load_workbook
import os
import numpy as np


app = Flask(__name__)

# 🔑 Twilio credentials (REPLACE THESE)
account_sid = "John doe"
auth_token = "Jane doe"

# 🧠 Tesseract path (Windows)
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


# 🔁 Download with retry (fixes timeout errors)
def download_image_with_retry(url, retries=3):
    for i in range(retries):
        try:
            response = requests.get(
                url,
                auth=HTTPBasicAuth(account_sid, auth_token),
                timeout=10
            )
            response.raise_for_status()
            return response.content

        except requests.exceptions.RequestException as e:
            print(f"Retry {i+1} failed:", e)
            time.sleep(2)

    return None

def extract_datetime(text):
    text = text.replace("\n", " ")

    date_pattern = (
        r"("
        r"(?:[A-Za-z]{3,9}\s*[-/,]?\s*\d{1,2}\s*[-/,]?\s*\d{4})"
        r"|"
        r"(?:\d{1,2}\s*[-/,]?\s*[A-Za-z]{3,9}\s*[-/,]?\s*\d{4})"
        r")"
    )

    time_pattern = r"(\d{1,2}:\d{2}(?::\d{2})?\s*[APMapm]{2}|\b\d{1,2}:\d{2}(?::\d{2})?\b)"

    date_match = re.search(date_pattern, text)
    time_match = re.search(time_pattern, text)

    date = date_match.group(1).strip() if date_match else ""
    time_val = time_match.group(1).strip() if time_match else ""

    return date, time_val

def extract_name(text):
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    return lines[-1] if lines else ""
def normalize_date(date_str):
    if not date_str:
        return ""

    date_str = re.sub(r"\s+", " ", date_str.strip())
    date_str = re.sub(r"\s*([-/,])\s*", r"\1", date_str)

    formats = [
        "%d %b %Y",
        "%d %B %Y",
        "%b %d %Y",
        "%B %d %Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%b-%d-%Y",
        "%B-%d-%Y",
        "%d/%b/%Y",
        "%d/%B/%Y",
        "%b/%d/%Y",
        "%B/%d/%Y",
        "%b %d, %Y",
        "%B %d, %Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(date_str, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue

    return date_str


def normalize_time(time_str):
    if not time_str:
        return ""

    time_str = re.sub(r"\s+", " ", time_str.strip()).upper()

    formats = [
        "%I:%M %p",
        "%I:%M:%S %p",
        "%H:%M",
        "%H:%M:%S",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(time_str, fmt).strftime("%H:%M")
        except ValueError:
            continue

    return time_str

def save_to_excel(date, time_val, name):
    file_name = "data.xlsx"

    # normalize before saving
    date = normalize_date(date)
    time_val = normalize_time(time_val)

    if not os.path.exists(file_name):
        wb = Workbook()
        ws = wb.active
        ws.append(["Date", "Time", "Name"])
    else:
        wb = load_workbook(file_name)
        ws = wb.active

    ws.append([date, time_val, name])
    wb.save(file_name)

    print("📊 Data saved to Excel")

def preprocess_image(input_path="received.jpg"):
    # Read image
    img = cv2.imread(input_path)

    # 1. CROP (bottom-right region)

    h, w = img.shape[:2]
    cropped = img[int(h * 0.6):h, int(w * 0.3):w]
    cv2.imwrite("debug_crop.jpg", cropped)

    # 2. Convert to HSV
    hsv = cv2.cvtColor(cropped, cv2.COLOR_BGR2HSV)

    # 3. Detect RED color

    lower_red1 = np.array([0, 120, 70])
    upper_red1 = np.array([10, 255, 255])

    lower_red2 = np.array([170, 120, 70])
    upper_red2 = np.array([180, 255, 255])

    mask1 = cv2.inRange(hsv, lower_red1, upper_red1)
    mask2 = cv2.inRange(hsv, lower_red2, upper_red2)

    mask = mask1 + mask2

    cv2.imwrite("debug_mask.jpg", mask)

    # 4. Apply mask (remove background)

    result = cv2.bitwise_and(cropped, cropped, mask=mask)
    cv2.imwrite("debug_result.jpg", result)

    # 5. Convert to grayscale

    gray = cv2.cvtColor(result, cv2.COLOR_BGR2GRAY)

    # 6. Threshold (make text white)

    _, thresh = cv2.threshold(gray, 50, 255, cv2.THRESH_BINARY)
    cv2.imwrite("debug_thresh.jpg", thresh)

    # 7. Clean noise

    kernel = np.ones((2, 2), np.uint8)
    clean = cv2.morphologyEx(thresh, cv2.MORPH_CLOSE, kernel)

    # 8. Make text thicker (important)

    clean = cv2.dilate(clean, kernel, iterations=1)
    cv2.imwrite("processed.jpg", clean)
    return "processed.jpg"

@app.route("/webhook", methods=["POST"])
def webhook():
    num_media = int(request.values.get("NumMedia", 0))

    if num_media > 0:
        image_url = request.values.get("MediaUrl0")
        print("📸 Image received:", image_url)

        img_data = download_image_with_retry(image_url)

        if img_data:
            # Save original image
            with open("received.jpg", "wb") as f:
                f.write(img_data)

            print("✅ Image saved")

            try:
                processed_path = preprocess_image("received.jpg")

                text = pytesseract.image_to_string(
                    processed_path,
                    config="--oem 3 --psm 6"
                        
                )

                print("\n📄 OCR Text:\n", text)

                # Extract required fields
                date, time_val = extract_datetime(text)
                name = extract_name(text)

                print("\n✅ Extracted Data:")
                print("Date:", date)
                print("Time:", time_val)
                print("Name:", name)

                # Save to Excel
                if date and time_val and name:
                    save_to_excel(date, time_val, name)
                else:
                    print("❌ Missing data, not saved")

            except Exception as e:
                print("❌ OCR error:", e)

        else:
            print("❌ Failed to download image after retries")

    else:
        incoming_msg = request.values.get("Body", "")
        print("💬 Text message:", incoming_msg)

    return "OK", 200


if __name__ == "__main__":
    app.run(port=5000, debug=True)